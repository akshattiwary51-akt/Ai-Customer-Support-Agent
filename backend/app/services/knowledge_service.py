import re

from fastapi import HTTPException, UploadFile, status
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.config import get_settings
from app.models import Document, DocumentChunk
from app.schemas import Citation

TOKEN_RE = re.compile(r"[a-zA-Z0-9]{3,}")
SUPPORTED_EXTENSIONS = {".txt", ".md", ".markdown", ".pdf", ".docx"}


def tokenize(text: str) -> set[str]:
    return {token.lower() for token in TOKEN_RE.findall(text)}


def chunk_text(text: str, size: int = 900, overlap: int = 140) -> list[str]:
    clean = re.sub(r"\s+", " ", text).strip()
    if not clean:
        return []
    chunks: list[str] = []
    start = 0
    while start < len(clean):
        chunk = clean[start : start + size].strip()
        if chunk:
            chunks.append(chunk)
        start += size - overlap
    return chunks


async def extract_text(file: UploadFile) -> tuple[str, str]:
    filename = file.filename or "upload.txt"
    lower = filename.lower()
    extension = "." + lower.rsplit(".", 1)[-1] if "." in lower else ""
    if extension not in SUPPORTED_EXTENSIONS:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Unsupported file type")
    data = await file.read()
    if len(data) > get_settings().max_upload_bytes:
        raise HTTPException(status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE, detail="File too large")

    if extension in {".txt", ".md", ".markdown"}:
        return data.decode("utf-8", errors="replace"), "text/markdown" if extension != ".txt" else "text/plain"
    if extension == ".pdf":
        try:
            from io import BytesIO

            from pypdf import PdfReader
        except ImportError as exc:
            raise HTTPException(
                status_code=status.HTTP_501_NOT_IMPLEMENTED,
                detail="PDF extraction requires the optional documents dependencies",
            ) from exc
        reader = PdfReader(BytesIO(data))
        return "\n".join(page.extract_text() or "" for page in reader.pages), "application/pdf"
    try:
        from io import BytesIO

        from docx import Document as DocxDocument
    except ImportError as exc:
        raise HTTPException(
            status_code=status.HTTP_501_NOT_IMPLEMENTED,
            detail="DOCX extraction requires the optional documents dependencies",
        ) from exc
    doc = DocxDocument(BytesIO(data))
    return "\n".join(paragraph.text for paragraph in doc.paragraphs), (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )


def index_document(db: Session, name: str, text: str, content_type: str, source: str = "upload") -> Document:
    document = Document(name=name, source=source, content_type=content_type, status="INDEXED")
    db.add(document)
    db.flush()
    chunks = chunk_text(text)
    if not chunks:
        document.status = "FAILED"
        document.error_message = "No extractable text found"
        db.commit()
        db.refresh(document)
        return document
    for index, chunk in enumerate(chunks):
        db.add(
            DocumentChunk(
                document_id=document.id,
                chunk_index=index,
                page_number=None,
                text=chunk,
                token_set=" ".join(sorted(tokenize(chunk))),
            )
        )
    db.commit()
    db.refresh(document)
    return document


def search_knowledge(db: Session, query: str, top_k: int | None = None) -> list[Citation]:
    settings = get_settings()
    query_tokens = tokenize(query)
    if not query_tokens:
        return []
    chunks = db.execute(select(DocumentChunk).join(Document)).scalars().all()
    ranked: list[tuple[float, DocumentChunk]] = []
    for chunk in chunks:
        tokens = set(chunk.token_set.split())
        if not tokens:
            continue
        overlap = len(query_tokens & tokens)
        score = overlap / max(len(query_tokens | tokens), 1)
        if score >= settings.retrieval_min_score:
            ranked.append((score, chunk))
    ranked.sort(key=lambda item: item[0], reverse=True)
    citations: list[Citation] = []
    for score, chunk in ranked[: top_k or settings.retrieval_top_k]:
        citations.append(
            Citation(
                document_id=chunk.document_id,
                document_name=chunk.document.name,
                chunk_id=chunk.id,
                page_number=chunk.page_number,
                score=round(score, 4),
                excerpt=chunk.text[:420],
            )
        )
    return citations
